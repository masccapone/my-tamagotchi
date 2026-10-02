"""거래내역 / 업체마스터 엑셀 템플릿을 생성한다. 실행: python3 tools/make_templates.py"""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

OUT = os.path.join(os.path.dirname(__file__), '..', 'templates')
HEAD = PatternFill('solid', fgColor='2A78D6')
OPT = PatternFill('solid', fgColor='8A8985')
EX = PatternFill('solid', fgColor='F1F0EC')
thin = Side(style='thin', color='D8D7D2')
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)


def header(ws, cols, widths, optional=()):
    for i, (c, w) in enumerate(zip(cols, widths), start=1):
        cell = ws.cell(row=1, column=i, value=c)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = OPT if c in optional else HEAD
        cell.alignment = Alignment(horizontal='center', vertical='center')
        ws.column_dimensions[cell.column_letter].width = w
    ws.freeze_panes = 'A2'


def guide(wb, title, lines):
    ws = wb.active
    ws.title = '작성안내'
    ws.column_dimensions['A'].width = 110
    ws['A1'] = title
    ws['A1'].font = Font(bold=True, size=14)
    for i, ln in enumerate(lines, start=3):
        ws.cell(row=i, column=1, value=ln).alignment = Alignment(wrap_text=True, vertical='top')
    return ws


def vendor_master(wb):
    ws = wb.create_sheet('업체마스터')
    header(ws, ['업체명(대표)', '별칭1', '별칭2', '별칭3', '결제기한(일)', '비고'], [26, 22, 22, 22, 14, 36],
           optional=('별칭1', '별칭2', '별칭3', '결제기한(일)', '비고'))
    ws.append(['(예시) 가나다환경', '가나다환경(주)', '주식회사 가나다환경', None, 30, '예시 행입니다. 지우고 사용하세요.'])
    for c in range(1, 7):
        ws.cell(row=2, column=c).fill = EX
    return ws


def trades():
    wb = Workbook()
    guide(wb, '거래내역 템플릿 작성 안내', [
        '■ 한 행에 거래 한 건을 적습니다. 처리비와 운반비처럼 금액이 둘이면 두 행으로 나누어 적습니다.',
        '■ 필수 항목: 일자, 방향, 거래처, 수량, 단위, 단가 (금액은 수량×단가로 자동 계산됩니다)',
        '■ 거래처는 [업체마스터] 시트에 등록된 이름을 목록에서 고릅니다. 새 업체는 먼저 업체마스터에 추가하세요. (표기가 달라지는 것을 막기 위해서입니다)',
        '■ 청구처: 실제 물량을 가져온(보내는) 업체와 대금을 청구·정산하는 업체가 다를 때만 적습니다. (예: 배출업체 A의 물량을 운송사 B가 모아서 정산하면 거래처=A, 청구처=B) 같으면 비워 두세요.',
        '■ 단가는 "수량 단위당 원"입니다. 단위가 kg이면 원/kg, 회이면 회당 원입니다.',
        '■ 금액이 수량×단가와 다르면 금액 칸에 직접 적으세요. (수식이 지워져도 괜찮습니다)',
        '■ 방향 이름(반입/반출 등)과 돈의 흐름은 [설정] 시트에서 바꿀 수 있습니다. 반입=받을 돈(매출), 반출=줄 돈(비용)으로 기본 설정되어 있습니다.',
        '■ 회색 예시 행은 삭제하고 사용하세요.',
        '■ 파일에는 거래처명과 금액이 들어 있으므로 외부에 공유하지 마세요.',
    ])
    ws = wb.create_sheet('거래내역')
    cols = ['일자', '방향', '거래처', '청구처', '품목', '수량', '단위', '단가종류', '단가', '금액', '비고']
    header(ws, cols, [12, 10, 26, 22, 14, 12, 8, 12, 12, 14, 30], optional=('청구처', '품목', '단가종류', '비고'))
    ex = [['2026-04-15', '반입', '(예시) 가나다환경', None, '폐합성수지류', 17570, 'kg', '처리비', 85, None, '예시 행'],
          ['2026-04-15', '반입', '(예시) 가나다환경', None, '폐합성수지류', 17570, 'kg', '운반비', 25, None, '예시 행'],
          ['2026-04-15', '반출', '(예시) 라마바시멘트', None, '20mm', 19550, 'kg', '처리비', 15, None, '예시 행']]
    f = lambda r: f'=IF(AND(F{r}<>"",I{r}<>""),F{r}*I{r},"")'
    for r, row in enumerate(ex, start=2):
        for c, v in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=v).fill = EX
        ws.cell(row=r, column=10, value=f(r)).fill = EX
    for r in range(5, 2001):
        ws.cell(row=r, column=10, value=f(r))
        ws.cell(row=r, column=1).number_format = 'yyyy-mm-dd'
    for r in range(2, 2001):
        ws.cell(row=r, column=6).number_format = '#,##0.##'
        ws.cell(row=r, column=9).number_format = '#,##0.##'
        ws.cell(row=r, column=10).number_format = '#,##0'
    for r in range(2, 5):
        ws.cell(row=r, column=1).number_format = '@'

    vendor_master(wb)
    st = wb.create_sheet('설정')
    header(st, ['방향 이름', '돈의 흐름', '', '단가종류', '', '단위'], [14, 14, 4, 14, 4, 10])
    for row in [('반입', '매출'), ('반출', '비용')]:
        st.append(list(row))
    for i, v in enumerate(['처리비', '운반비', '기타'], start=2):
        st.cell(row=i, column=4, value=v)
    for i, v in enumerate(['kg', '톤', '회', '건'], start=2):
        st.cell(row=i, column=6, value=v)
    st['H1'] = '※ 돈의 흐름은 "매출"(받을 돈) 또는 "비용"(줄 돈)만 쓸 수 있습니다.'

    dv_dir = DataValidation(type='list', formula1='=설정!$A$2:$A$20', allow_blank=True)
    dv_vendor = DataValidation(type='list', formula1='=업체마스터!$A$2:$A$500', allow_blank=True, showErrorMessage=False)
    dv_bill = DataValidation(type='list', formula1='=업체마스터!$A$2:$A$500', allow_blank=True, showErrorMessage=False)
    dv_unit = DataValidation(type='list', formula1='=설정!$F$2:$F$20', allow_blank=True)
    dv_kind = DataValidation(type='list', formula1='=설정!$D$2:$D$20', allow_blank=True)
    for dv, col in [(dv_dir, 'B'), (dv_vendor, 'C'), (dv_bill, 'D'), (dv_unit, 'G'), (dv_kind, 'H')]:
        ws.add_data_validation(dv)
        dv.add(f'{col}2:{col}2000')
    wb.move_sheet('거래내역', offset=-1)
    wb.active = 0
    wb.save(os.path.join(OUT, '거래내역_템플릿.xlsx'))


def vendors_only():
    wb = Workbook()
    guide(wb, '업체마스터 작성 안내', [
        '■ 업체명(대표): 화면과 리포트에 표시할 이름 하나만 정합니다.',
        '■ 별칭: ERP(전표), 은행 입금자명, 현장 장부에서 같은 업체를 다르게 부르는 이름을 적습니다. 이름이 연결되어야 "A업체 미수금"을 정확히 찾습니다.',
        '■ 결제기한(일): 아는 업체만 적습니다. 청구일로부터 며칠 뒤 결제하는지입니다. 비워 두면 과거 결제 이력의 평균으로 추정합니다.',
        '■ 신규 거래처가 생길 때만 추가하면 됩니다.',
    ])
    vendor_master(wb)
    wb.save(os.path.join(OUT, '업체마스터_템플릿.xlsx'))


if __name__ == '__main__':
    os.makedirs(OUT, exist_ok=True)
    trades()
    vendors_only()
    print('ok')
